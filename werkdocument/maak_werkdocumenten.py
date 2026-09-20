#!/usr/bin/env python3
"""
maak_werkdocumenten.py
Genereert de bestanden voor les 2 "Mijn digitaal stageportfolio" (De Speelboom):

  - Portfoliopaspoort.docx              het werkdocument dat de leerling INLEVERT
  - Document zonder titel.docx          rommelbestand 1 (samen hernoemd tijdens de demo)
  - verslag Liam boos DEFINITIEF (2).docx   rommelbestand 2
  - dingen.xlsx                         rommelbestand 3
  - Kopie van sjabloon reflectie.docx    rommelbestand 4
  - uurrooster okt def.docx             rommelbestand 5

Upload ze naar Google Drive en open ze met Google Documenten / Google Spreadsheets.

Gebruik:  python3 maak_werkdocumenten.py     (vereist python-docx en openpyxl)

LET OP bij aanpassen:
 - De slechte bestandsnamen zijn OPZETTELIJK. Ze zijn het lesmateriaal.
 - In elk rommelbestand staat bovenaan een datum in de vorm DD-MM-JJJJ.
   De leerling moet die omzetten naar JJJJ-MM-DD. Wijzig die vorm dus niet.
 - Alle namen, adressen en gegevens zijn fictief.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font as XFont, PatternFill, Alignment as XAlign

NAVY = RGBColor(0x2A, 0x39, 0x73)
GREY = RGBColor(0x4D, 0x55, 0x73)
HERE = os.path.dirname(os.path.abspath(__file__))


# ---------- hulpfuncties ----------
def basis_document(marge=2.0):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(marge)
        s.left_margin = s.right_margin = Cm(marge)
    return doc


def kop(doc, tekst, grootte=16, ruimte_voor=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(ruimte_voor)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(tekst)
    r.bold = True
    r.font.size = Pt(grootte)
    r.font.color.rgb = NAVY
    return p


def tekst(doc, s, cursief=False, klein=False, vet=False, na=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(na)
    r = p.add_run(s)
    r.italic = cursief
    r.bold = vet
    if klein:
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY
    return p


def schaduw(cel, kleur="E4E8F6"):
    tcPr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), kleur)
    tcPr.append(shd)


def antwoordlijnen(doc, aantal=2):
    for _ in range(aantal):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run("_" * 86)


# ---------- 1. Portfoliopaspoort (het werkdocument) ----------
def portfoliopaspoort():
    doc = basis_document()

    p = doc.add_paragraph()
    r = p.add_run("PORTFOLIOPASPOORT")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    tekst(doc, "Les 2 — Mijn digitaal stageportfolio  ·  Toegepaste Informatica  ·  De Speelboom",
          klein=True, na=10)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, label in enumerate(["Naam:", "Klas:", "Datum:"]):
        c = t.rows[0].cells[i]
        c.text = label + " "
        c.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    kop(doc, "Zo werk je", 13, ruimte_voor=6)
    for s in [
        "1.  Links op je scherm staat de lespagina. Daar lees je wat je moet doen.",
        "2.  In Google Drive bouw je je mappen. Dat is je echte werk.",
        "3.  In dit document schrijf je op wat je gedaan hebt.",
        "4.  Alleen DIT document lever je in via Google Classroom.",
    ]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(2)

    # ---- Deel 1
    kop(doc, "Deel 1 — Waar staat mijn bestand?")
    tekst(doc, "Vraag 1. Je bewaart een verslag in de map Downloads van je Chromebook. Morgen krijg je een "
               "ander toestel. Kan je je verslag dan nog openen? Leg uit in één zin.", vet=True)
    antwoordlijnen(doc, 2)

    # ---- Deel 2
    kop(doc, "Deel 2 — Mijn vijf bestanden")
    tekst(doc, "Vul de tabel in terwijl je werkt. Het eerste bestand deed je samen met je leraar.", klein=True)

    t = doc.add_table(rows=6, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["Oude naam", "Mijn nieuwe naam", "In welke map?"]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        schaduw(c)
    rijen = [
        ("Document zonder titel", "2026-09-22_activiteitenfiche_herfstslinger_v1", "02_Activiteiten"),
        ("verslag Liam boos DEFINITIEF (2)", "", ""),
        ("dingen", "", ""),
        ("Kopie van sjabloon reflectie", "", ""),
        ("uurrooster okt def", "", ""),
    ]
    for i, (oud, nieuw, mp) in enumerate(rijen, start=1):
        cells = t.rows[i].cells
        cells[0].text = oud
        cells[1].text = nieuw
        cells[2].text = mp
        for c in cells:
            c.paragraphs[0].runs and setattr(c.paragraphs[0].runs[0].font, "size", Pt(10))
    doc.add_paragraph()

    tekst(doc, "Vraag 2. Waarom zetten we de datum vóóraan in de bestandsnaam, en waarom in de vorm "
               "2026-09-22?", vet=True)
    antwoordlijnen(doc, 2)
    tekst(doc, "Vraag 3. Waarom mag de naam van een kind niet in een bestandsnaam staan?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- Deel 3
    kop(doc, "Deel 3 — Een foto van mijn kast")
    tekst(doc, "Plak hieronder je screenshot: Invoegen › Afbeelding › Uploaden vanaf computer › Downloads.",
          klein=True)
    tekst(doc, "Op je screenshot moet je de naam van je hoofdmap én je vijf submappen kunnen lezen.", klein=True)
    for _ in range(6):
        doc.add_paragraph()

    # ---- Deel 4
    kop(doc, "Deel 4 — Delen met mijn leraar")
    tekst(doc, "Vraag 4. Je deelde je map met je leraar als Lezer. Noem één ding dat je leraar nu WÉL mag "
               "doen en één ding dat je leraar NIET mag doen.", vet=True)
    tekst(doc, "Mijn leraar mag wel:")
    antwoordlijnen(doc, 1)
    tekst(doc, "Mijn leraar mag niet:")
    antwoordlijnen(doc, 1)

    # ---- Deel 5
    kop(doc, "Deel 5 — Slotvragen")
    tekst(doc, "Vraag 5. Noem één ding dat jij vanaf nu anders gaat doen met je bestanden.", vet=True)
    antwoordlijnen(doc, 1)
    tekst(doc, "Vraag 6. Welke stap van vandaag was voor jou het moeilijkst? Waarom?", vet=True)
    antwoordlijnen(doc, 2)

    # ---- Zelfcontrole
    kop(doc, "Zelfcontrole — aankruisen vóór je indient")
    for s in [
        "Mijn hoofdmap staat in Mijn Drive en heeft mijn naam.",
        "Ik heb 5 submappen, genummerd 01 tot 05.",
        "Mijn 5 bestanden hebben een nieuwe naam volgens de naamafspraak.",
        "In geen enkele bestandsnaam staat de naam van een kind.",
        "Elk bestand staat in de juiste map.",
        "Mijn screenshot staat in deel 3.",
        "Mijn map is gedeeld met mijn leraar als Lezer.",
        "Alle vragen zijn ingevuld.",
    ]:
        p = doc.add_paragraph("☐  " + s)
        p.paragraph_format.space_after = Pt(2)

    # ---- Extra
    kop(doc, "Optionele uitbreiding")
    tekst(doc, "Alleen als al de rest af is en ingeleverd. Welke zesde map heb jij nodig op jouw stageplaats, "
               "en waarom?", klein=True)
    antwoordlijnen(doc, 2)

    doc.save(os.path.join(HERE, "Portfoliopaspoort.docx"))


# ---------- 2. De rommelbestanden ----------
def rommel_activiteitenfiche():
    doc = basis_document(marge=2.5)
    tekst(doc, "Datum: 22-09-2026", vet=True)
    tekst(doc, "Groep: 6 tot 8 jaar — 10 kinderen")
    doc.add_paragraph()
    tekst(doc, "Knutselen: herfstslinger", vet=True)
    tekst(doc, "Doel: de kinderen knippen en rijgen zelf een slinger van herfstbladeren voor het raam "
               "van de leefruimte.")
    tekst(doc, "Materiaal: gekleurd papier, schaar, touw, perforator, bladeren van de speelplaats.")
    tekst(doc, "Verloop:")
    for s in ["1. Bladeren zoeken op de speelplaats (10 minuten).",
              "2. Bladeren op papier tekenen en uitknippen (15 minuten).",
              "3. Gaatjes maken en aan het touw rijgen (15 minuten).",
              "4. Samen ophangen en opruimen (10 minuten)."]:
        p = doc.add_paragraph(s)
        p.paragraph_format.space_after = Pt(2)
    tekst(doc, "Aandacht voor veiligheid: scharen met ronde punt, perforator alleen samen met de begeleider.")
    doc.save(os.path.join(HERE, "Document zonder titel.docx"))


def rommel_observatie():
    doc = basis_document(marge=2.5)
    tekst(doc, "Datum: 18-09-2026", vet=True)
    tekst(doc, "Wat ik zag in de opvang", vet=True)
    doc.add_paragraph()
    tekst(doc, "Liam (4 jaar) speelde in de bouwhoek. Een ander kind nam een blok van zijn toren. "
               "Liam riep luid “neen” en duwde het kind weg. De toren viel om.")
    tekst(doc, "Ik ging bij hem zitten. Na ongeveer twee minuten begon hij opnieuw te bouwen, "
               "eerst alleen, daarna samen met hetzelfde kind.")
    tekst(doc, "Dit is mijn tweede versie. In de eerste versie schreef ik “Liam was stout”. "
               "Dat is geen feit maar een mening, dus heb ik het aangepast.")
    doc.save(os.path.join(HERE, "verslag Liam boos DEFINITIEF (2).docx"))


def rommel_reflectie():
    doc = basis_document(marge=2.5)
    tekst(doc, "Datum: 21-09-2026", vet=True)
    tekst(doc, "Reflectie week 1", vet=True)
    doc.add_paragraph()
    tekst(doc, "Wat ging goed?")
    antwoordlijnen(doc, 2)
    tekst(doc, "Wat was moeilijk?")
    antwoordlijnen(doc, 2)
    tekst(doc, "Wat doe ik volgende week anders?")
    antwoordlijnen(doc, 2)
    doc.save(os.path.join(HERE, "Kopie van sjabloon reflectie.docx"))


def rommel_uurrooster():
    doc = basis_document(marge=2.5)
    tekst(doc, "Datum: 01-10-2026", vet=True)
    tekst(doc, "Mijn uurrooster op De Speelboom — oktober", vet=True)
    doc.add_paragraph()
    t = doc.add_table(rows=5, cols=3)
    t.style = "Table Grid"
    rijen = [("Dag", "Van", "Tot"),
             ("maandag", "15.30", "18.00"),
             ("woensdag", "12.00", "17.00"),
             ("donderdag", "15.30", "18.00"),
             ("vrijdag", "15.30", "17.30")]
    for i, rij in enumerate(rijen):
        for j, v in enumerate(rij):
            c = t.rows[i].cells[j]
            c.text = v
            if i == 0:
                c.paragraphs[0].runs[0].bold = True
                schaduw(c)
    doc.add_paragraph()
    tekst(doc, "Contact: onthaal De Speelboom, 09 000 00 00 (fictief nummer).", klein=True)
    doc.save(os.path.join(HERE, "uurrooster okt def.docx"))


def rommel_boodschappen():
    wb = Workbook()
    ws = wb.active
    ws.title = "Blad1"
    ws["A1"] = "Datum: 22-09-2026"
    ws["A1"].font = XFont(name="Arial", bold=True)
    ws["A2"] = "Nodig voor het knutselmoment van dinsdag"
    ws["A2"].font = XFont(name="Arial", italic=True)
    kop = ["Wat", "Hoeveel", "Waar"]
    for i, h in enumerate(kop, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = XFont(name="Arial", bold=True)
        c.fill = PatternFill("solid", fgColor="E4E8F6")
    data = [("gekleurd papier A4", "2 pakken", "kast bergruimte"),
            ("touw", "1 rol", "nog kopen"),
            ("perforator", "3 stuks", "kast bergruimte"),
            ("scharen (ronde punt)", "10 stuks", "leefruimte"),
            ("koekjes voor het vieruurtje", "1 doos", "nog kopen")]
    for r, rij in enumerate(data, start=5):
        for c, v in enumerate(rij, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = XFont(name="Arial")
            cell.alignment = XAlign(vertical="center")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 22
    wb.save(os.path.join(HERE, "dingen.xlsx"))


if __name__ == "__main__":
    portfoliopaspoort()
    rommel_activiteitenfiche()
    rommel_observatie()
    rommel_reflectie()
    rommel_uurrooster()
    rommel_boodschappen()
    print("Klaar. Bestanden staan in:", HERE)
